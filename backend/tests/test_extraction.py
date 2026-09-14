import os
import io
import pandas as pd
from docx import Document
from pypdf import PdfWriter, PdfReader

from app.services.extraction.pdf_extractor import PDFExtractor
from app.services.extraction.docx_extractor import DOCXExtractor
from app.services.extraction.xlsx_extractor import XLSXExtractor

def test_pdf_extraction(tmp_path):
    # Create a dummy PDF
    pdf_path = tmp_path / "test.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with open(pdf_path, "wb") as f:
        writer.write(f)

    # Test
    extractor = PDFExtractor()
    elements = list(extractor.extract(str(pdf_path), "test.pdf"))
    
    # Since it's a blank page, pypdf extract_text() returns empty string and we skip empty pages.
    assert len(elements) == 0

def test_docx_extraction(tmp_path):
    docx_path = tmp_path / "test.docx"
    doc = Document()
    doc.add_paragraph("Hello world")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A1"
    table.cell(0, 1).text = "B1"
    table.cell(1, 0).text = "A2"
    table.cell(1, 1).text = "B2"
    doc.save(docx_path)

    extractor = DOCXExtractor()
    elements = list(extractor.extract(str(docx_path), "test.docx"))

    assert len(elements) == 2
    assert elements[0].element_type == "text"
    assert elements[0].content == "Hello world"
    assert elements[0].metadata_["type"] == "paragraph"

    assert elements[1].element_type == "table"
    assert "A1 | B1" in elements[1].content
    assert "A2 | B2" in elements[1].content
    assert elements[1].metadata_["type"] == "table"

def test_xlsx_extraction(tmp_path):
    xlsx_path = tmp_path / "test.xlsx"
    df = pd.DataFrame({"A": [1, 2], "B": [3, 4]})
    df.to_excel(xlsx_path, index=False, sheet_name="Sheet1")

    extractor = XLSXExtractor()
    elements = list(extractor.extract(str(xlsx_path), "test.xlsx"))

    assert len(elements) == 1
    assert elements[0].element_type == "table"
    assert "Sheet1" == elements[0].metadata_["sheet_name"]
    assert "1" in elements[0].content
    assert "3" in elements[0].content
